"""One-off script to generate the PDF and DOCX test fixtures.

Not part of the application - run once to produce
tests/fixtures/sample.pdf and tests/fixtures/sample.docx from the same
fictional content used in sample.md / sample.html. Safe to re-run any time
these fixtures need regenerating.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures"

TITLE = "Post-Operative Care Guidelines: Total Knee Replacement"
DISCLAIMER = "Fictional sample document for testing purposes only. Not real medical advice."

SECTIONS = [
    (
        "Overview",
        "This guideline outlines standard post-operative care recommendations for "
        "patients recovering from total knee replacement surgery (also known as total "
        "knee arthroplasty). It is intended for use by nursing staff and physical "
        "therapists during the first six weeks of recovery.",
    ),
    (
        "Pain Management",
        "Patients typically receive a combination of oral analgesics and ice therapy "
        "during the first 72 hours. Pain should be reassessed every four hours using a "
        "standard 0-10 pain scale. Escalating pain accompanied by swelling or warmth at "
        "the incision site should be reported to the attending physician immediately, "
        "as this may indicate infection or deep vein thrombosis.",
    ),
    (
        "Mobility and Physical Therapy",
        "Weight-bearing as tolerated is generally encouraged starting on the first "
        "post-operative day. A physical therapist should evaluate range of motion daily "
        "for the first week. Most patients achieve 90 degrees of knee flexion within two "
        "weeks and should be walking with a cane or no assistive device by six weeks, "
        "depending on baseline fitness and adherence to the exercise program.",
    ),
    (
        "Wound Care",
        "The surgical dressing should remain dry and intact for the first 48 hours. "
        "Sutures or staples are typically removed between day 10 and day 14. Patients "
        "should be instructed to watch for redness, discharge, or fever, all of which "
        "warrant a follow-up visit.",
    ),
    (
        "Follow-Up Schedule",
        "A follow-up appointment is recommended at two weeks, six weeks, and three "
        "months post-surgery to assess healing, range of motion, and overall functional "
        "recovery.",
    ),
]


def build_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER)
    story = [Paragraph(TITLE, styles["Title"]), Paragraph(DISCLAIMER, styles["Italic"]), Spacer(1, 12)]
    for heading, body in SECTIONS:
        story.append(Paragraph(heading, styles["Heading2"]))
        story.append(Paragraph(body, styles["BodyText"]))
        story.append(Spacer(1, 8))
    doc.build(story)


def build_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(DISCLAIMER).italic = True
    for heading, body in SECTIONS:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
    doc.save(str(path))


if __name__ == "__main__":
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    build_pdf(FIXTURES_DIR / "sample.pdf")
    build_docx(FIXTURES_DIR / "sample.docx")
    print(f"Wrote sample.pdf and sample.docx to {FIXTURES_DIR}")
